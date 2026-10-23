"""Baseline mapping store (the only persisted state).

Confirmed mappings are written per BMC per domain as an Excel file (on ADLS in
production). Columns mirror the notebook's mapping sheet plus a weight encoding
so splits round-trip:

    raw_var_name | raw_region | processed_var_name | processed_region |
    percent | action | weight_kind | weight_value

On the next refresh the baseline is loaded, keyed by source, and fed to the
resolver so stable rows carry over and only deltas need review.
"""
from __future__ import annotations

import json
import os
from pathlib import Path

import numpy as np

from .actions import TemporalBreakpoint, Weight
from .regions import RegionKey
from .resolver import MappingProposal, TargetRef

COLUMNS = ["raw_var_name", "raw_region", "processed_var_name", "processed_region",
           "percent", "action", "weight_kind", "weight_value"]


def _weight_to_cells(w: Weight) -> tuple[str, str]:
    if w.kind == "scalar":
        return "scalar", repr(float(w.value))
    if w.kind == "vector":
        return "vector", json.dumps(np.asarray(w.value, float).tolist())
    if w.kind == "rule":
        r = w.value
        return "rule", json.dumps({"breakpoint": r.breakpoint, "side": r.side})
    raise ValueError(w.kind)


def _weight_from_cells(kind: str, value: str) -> Weight:
    if kind == "scalar":
        return Weight.scalar(float(value))
    if kind == "vector":
        return Weight.vector(json.loads(value))
    if kind == "rule":
        d = json.loads(value)
        return Weight.rule(TemporalBreakpoint(d["breakpoint"], d.get("side", "after")))
    return Weight.scalar(1.0)


def proposals_to_rows(proposals: list[MappingProposal]) -> list[dict]:
    rows: list[dict] = []
    for p in proposals:
        if p.action == "unmapped" or not p.targets:
            continue
        if p.action == "aggregate":
            t = p.targets[0]
            for s in p.sources:
                k, v = _weight_to_cells(Weight.scalar(1.0))
                rows.append(_row(s, (t.var, t.region), t.percent, p.action, k, v))
        else:
            s = p.sources[0]
            for t in p.targets:
                k, v = _weight_to_cells(t.weight)
                rows.append(_row(s, (t.var, t.region), t.percent, p.action, k, v))
    return rows


def _row(src, tgt, percent, action, wkind, wval) -> dict:
    sv, sr = src
    tv, tr = tgt
    return {
        "raw_var_name": sv, "raw_region": RegionKey.parse(sr).canonical,
        "processed_var_name": tv, "processed_region": RegionKey.parse(tr).canonical,
        "percent": percent, "action": action, "weight_kind": wkind, "weight_value": wval,
    }


def rows_to_proposals(rows: list[dict]) -> list[MappingProposal]:
    singles: list[MappingProposal] = []
    splits: dict[tuple[str, str], list[dict]] = {}
    aggs: dict[tuple[str, str], list[dict]] = {}
    for r in rows:
        action = str(r["action"])
        if action == "split":
            splits.setdefault((r["raw_var_name"], r["raw_region"]), []).append(r)
        elif action == "aggregate":
            aggs.setdefault((r["processed_var_name"], r["processed_region"]), []).append(r)
        else:
            singles.append(_single(r, action))

    out = list(singles)
    for (sv, sr), group in splits.items():
        src = (sv, RegionKey.parse(sr))
        targets = [TargetRef(g["processed_var_name"], RegionKey.parse(g["processed_region"]),
                             _weight_from_cells(g["weight_kind"], str(g["weight_value"])),
                             float(g["percent"] or 0)) for g in group]
        out.append(_confirmed([src], "split", targets))
    for (tv, tr), group in aggs.items():
        sources = [(g["raw_var_name"], RegionKey.parse(g["raw_region"])) for g in group]
        tgt = [TargetRef(tv, RegionKey.parse(tr), Weight.scalar(1.0), 100.0)]
        out.append(_confirmed(sources, "aggregate", tgt))
    return out


def _single(r, action) -> MappingProposal:
    src = (r["raw_var_name"], RegionKey.parse(r["raw_region"]))
    tgt = [TargetRef(r["processed_var_name"], RegionKey.parse(r["processed_region"]),
                     _weight_from_cells(r["weight_kind"], str(r["weight_value"])),
                     float(r["percent"] or 0))]
    return _confirmed([src], action, tgt)


def _confirmed(sources, action, targets) -> MappingProposal:
    return MappingProposal(sources=sources, action=action, targets=targets,
                           confidence=1.0, provenance="baseline", status="confirmed")


class ExcelBaselineStore:
    """Baseline Excel files under base_dir/{bmc}/{domain}.xlsx (local or mounted ADLS)."""

    def __init__(self, base_dir: str):
        self.base = Path(base_dir)

    def _path(self, bmc: str, domain: str) -> Path:
        return self.base / bmc / f"{domain}.xlsx"

    def load(self, bmc: str, domain: str) -> list[MappingProposal]:
        path = self._path(bmc, domain)
        if not path.exists():
            return []
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = [str(h) for h in next(it, [])]
        rows = [dict(zip(header, r)) for r in it if any(c is not None for c in r)]
        return rows_to_proposals(rows)

    def save(self, bmc: str, domain: str, proposals: list[MappingProposal]) -> str:
        from openpyxl import Workbook
        path = self._path(bmc, domain)
        os.makedirs(path.parent, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "Mapping"
        ws.append(COLUMNS)
        for row in proposals_to_rows(proposals):
            ws.append([row[c] for c in COLUMNS])
        wb.save(path)
        return str(path)

    def baseline_dict(self, bmc: str, domain: str) -> dict:
        """Keyed by first source (var, RegionKey) for the resolver."""
        return {p.sources[0]: p for p in self.load(bmc, domain) if p.sources}


class NoBaselineStore:
    def load(self, bmc, domain): return []
    def save(self, bmc, domain, proposals): return ""
    def baseline_dict(self, bmc, domain): return {}
